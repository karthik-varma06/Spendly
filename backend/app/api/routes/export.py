from __future__ import annotations

from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO
from numbers import Number
from typing import Any

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.api.deps import get_current_user, get_db
from app.models.user import User

router = APIRouter(prefix="/export", tags=["export"])


# ---------------------------------------------------------------------
# Formatting helpers
# ---------------------------------------------------------------------
def _today() -> date:
    return date.today()


def _fmt_date(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat(timespec="seconds")
    return str(value)


def _fmt_bool(value: Any) -> str:
    return "Yes" if bool(value) else "No"


def _fmt_money_pdf(value: Any) -> str:
    try:
        return f"INR {float(value or 0):,.2f}"
    except Exception:
        return "INR 0.00"


def _fmt_money_excel(value: Any) -> float:
    try:
        return float(value or 0)
    except Exception:
        return 0.0


def _fmt_number_excel(value: Any) -> float:
    try:
        return float(value or 0)
    except Exception:
        return 0.0


def _fmt_money_csv(value: Any) -> str:
    try:
        return f"{float(value or 0):.2f}"
    except Exception:
        return "0.00"


def _fmt_number(value: Any) -> str:
    try:
        return f"{float(value or 0):.2f}"
    except Exception:
        return "0.00"


def _safe_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return _fmt_date(value)
    if isinstance(value, bool):
        return _fmt_bool(value)
    return str(value)


def _slugify_filename(name: str) -> str:
    chars = []
    for ch in (name or "spendly").strip().lower():
        if ch.isalnum():
            chars.append(ch)
        elif ch in {" ", "-", "_"}:
            chars.append("_")
    slug = "".join(chars).strip("_")
    return slug or "spendly"


# ---------------------------------------------------------------------
# DB helpers
# ---------------------------------------------------------------------
def _scalar(db: Session, sql: str, **params: Any) -> float:
    val = db.execute(text(sql), params).scalar()
    try:
        return float(val or 0)
    except Exception:
        return 0.0


def _rows(db: Session, sql: str, **params: Any) -> list[dict[str, Any]]:
    return [dict(r) for r in db.execute(text(sql), params).mappings().all()]


def _build_data(db: Session, user_id: str) -> dict[str, Any]:
    profile = _rows(
        db,
        """
        SELECT full_name, email, currency, timezone
        FROM users
        WHERE id = :user_id
        LIMIT 1
        """,
        user_id=user_id,
    )

    monthly_expense = _scalar(
        db,
        """
        SELECT COALESCE(SUM(total_amount), 0)
        FROM expenses
        WHERE user_id = :user_id
          AND expense_date >= date_trunc('month', CURRENT_DATE)::date
        """,
        user_id=user_id,
    )
    yearly_expense = _scalar(
        db,
        """
        SELECT COALESCE(SUM(total_amount), 0)
        FROM expenses
        WHERE user_id = :user_id
          AND expense_date >= date_trunc('year', CURRENT_DATE)::date
        """,
        user_id=user_id,
    )
    monthly_income = _scalar(
        db,
        """
        SELECT COALESCE(SUM(amount), 0)
        FROM income
        WHERE user_id = :user_id
          AND income_date >= date_trunc('month', CURRENT_DATE)::date
        """,
        user_id=user_id,
    )
    yearly_income = _scalar(
        db,
        """
        SELECT COALESCE(SUM(amount), 0)
        FROM income
        WHERE user_id = :user_id
          AND income_date >= date_trunc('year', CURRENT_DATE)::date
        """,
        user_id=user_id,
    )

    dashboard = {
        "monthly_expense": monthly_expense,
        "yearly_expense": yearly_expense,
        "monthly_income": monthly_income,
        "yearly_income": yearly_income,
        "monthly_savings": monthly_income - monthly_expense,
        "yearly_savings": yearly_income - yearly_expense,
        "lifetime_savings": _scalar(
            db,
            """
            SELECT COALESCE((SELECT SUM(amount) FROM income WHERE user_id = :user_id), 0)
                 - COALESCE((SELECT SUM(total_amount) FROM expenses WHERE user_id = :user_id), 0)
            """,
            user_id=user_id,
        ),
        "expense_count": _scalar(db, "SELECT COUNT(*) FROM expenses WHERE user_id = :user_id", user_id=user_id),
        "budget_count": _scalar(db, "SELECT COUNT(*) FROM budgets WHERE user_id = :user_id", user_id=user_id),
        "uploaded_count": _scalar(db, "SELECT COUNT(*) FROM uploaded_files WHERE user_id = :user_id", user_id=user_id),
        "top_categories": _rows(
            db,
            """
            SELECT COALESCE(c.name, 'OTHER') AS name,
                   COALESCE(SUM(e.total_amount), 0) AS amount
            FROM expenses e
            LEFT JOIN categories c ON c.id = e.category_id
            WHERE e.user_id = :user_id
            GROUP BY COALESCE(c.name, 'OTHER')
            ORDER BY amount DESC, name ASC
            LIMIT 10
            """,
            user_id=user_id,
        ),
        "top_merchants": _rows(
            db,
            """
            SELECT COALESCE(NULLIF(TRIM(vendor_name), ''), 'Unknown') AS name,
                   COALESCE(SUM(total_amount), 0) AS amount
            FROM expenses
            WHERE user_id = :user_id
            GROUP BY COALESCE(NULLIF(TRIM(vendor_name), ''), 'Unknown')
            ORDER BY amount DESC, name ASC
            LIMIT 10
            """,
            user_id=user_id,
        ),
        "top_items": _rows(
            db,
            """
            SELECT COALESCE(NULLIF(TRIM(ei.item_name), ''), 'Unknown') AS name,
                   COALESCE(SUM(ei.total_price), 0) AS amount
            FROM expense_items ei
            JOIN expenses e ON e.id = ei.expense_id
            WHERE e.user_id = :user_id
            GROUP BY COALESCE(NULLIF(TRIM(ei.item_name), ''), 'Unknown')
            ORDER BY amount DESC, name ASC
            LIMIT 10
            """,
            user_id=user_id,
        ),
        "spending_trends": _rows(
            db,
            """
            SELECT TO_CHAR(DATE_TRUNC('month', expense_date), 'Mon YYYY') AS month,
                   COALESCE(SUM(total_amount), 0) AS amount
            FROM expenses
            WHERE user_id = :user_id
              AND expense_date IS NOT NULL
            GROUP BY DATE_TRUNC('month', expense_date)
            ORDER BY DATE_TRUNC('month', expense_date)
            """,
            user_id=user_id,
        ),
    }

    budgets = _rows(
        db,
        """
        SELECT COALESCE(c.name, 'ALL') AS category_name,
               b.budget_limit,
               b.spent_amount,
               b.period,
               b.start_date,
               b.end_date,
               b.notify_percentage,
               b.is_active,
               b.created_at
        FROM budgets b
        LEFT JOIN categories c ON c.id = b.category_id
        WHERE b.user_id = :user_id
        ORDER BY b.created_at DESC
        """,
        user_id=user_id,
    )

    expenses = _rows(
        db,
        """
        SELECT COALESCE(e.vendor_name, '') AS vendor_name,
               COALESCE(c.name, 'OTHER') AS category_name,
               e.expense_date,
               e.expense_time,
               e.subtotal,
               e.cgst,
               e.sgst,
               e.igst,
               e.tax_amount,
               e.discount_amount,
               e.total_amount,
               e.payment_method,
               e.status,
               e.is_reviewed,
               e.is_manual,
               e.created_at AS created_at
        FROM expenses e
        LEFT JOIN categories c ON c.id = e.category_id
        WHERE e.user_id = :user_id
        ORDER BY e.expense_date DESC NULLS LAST, e.created_at DESC
        """,
        user_id=user_id,
    )

    expense_items = _rows(
        db,
        """
        SELECT COALESCE(NULLIF(TRIM(ei.item_name), ''), 'Unknown') AS item_name,
               COALESCE(ei.item_category, 'OTHER') AS item_category,
               ei.quantity,
               ei.unit_price,
               ei.total_price,
               COALESCE(e.vendor_name, 'Unknown') AS vendor_name,
               e.expense_date,
               ei.created_at
        FROM expense_items ei
        JOIN expenses e ON e.id = ei.expense_id
        WHERE e.user_id = :user_id
        ORDER BY e.expense_date DESC NULLS LAST, ei.created_at DESC
        """,
        user_id=user_id,
    )

    uploads = _rows(
        db,
        """
        SELECT
            original_filename,
            file_type,
            file_size,
            extraction_status,
            created_at
        FROM uploaded_files
        WHERE user_id = :user_id
        ORDER BY created_at DESC NULLS LAST
        """,
        user_id=user_id,
    )

    income = _rows(
        db,
        """
        SELECT COALESCE(source, 'Unknown') AS source,
               amount,
               income_date,
               COALESCE(notes, '') AS notes,
               created_at
        FROM income
        WHERE user_id = :user_id
        ORDER BY income_date DESC NULLS LAST, created_at DESC
        """,
        user_id=user_id,
    )

    return {
        "profile": profile,
        "dashboard": dashboard,
        "budgets": budgets,
        "expenses": expenses,
        "expense_items": expense_items,
        "uploads": uploads,
        "income": income,
    }


# ---------------------------------------------------------------------
# PDF helpers
# ---------------------------------------------------------------------
def _pdf_styles():
    styles = getSampleStyleSheet()

    styles.add(
        ParagraphStyle(
            name="SpendlyTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=22,
            leading=26,
            textColor=colors.HexColor("#0f172a"),
            alignment=TA_LEFT,
            spaceAfter=6,
        )
    )
    styles.add(
        ParagraphStyle(
            name="SpendlySubtle",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=8.5,
            leading=10,
            textColor=colors.HexColor("#475569"),
            spaceAfter=8,
        )
    )
    styles.add(
        ParagraphStyle(
            name="SpendlySection",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=12.5,
            leading=15,
            textColor=colors.HexColor("#111827"),
            spaceBefore=10,
            spaceAfter=6,
        )
    )
    styles.add(
        ParagraphStyle(
            name="SpendlyCell",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=7.8,
            leading=9.4,
            textColor=colors.HexColor("#0f172a"),
        )
    )
    styles.add(
        ParagraphStyle(
            name="SpendlyHeader",
            parent=styles["BodyText"],
            fontName="Helvetica-Bold",
            fontSize=7.9,
            leading=9.2,
            textColor=colors.white,
        )
    )
    return styles


def _page_number(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 9)
    canvas.setFillColor(colors.HexColor("#64748b"))
    canvas.drawRightString(doc.pagesize[0] - 18, 12, f"Page {doc.page}")
    canvas.restoreState()


def _pdf_table_from_dicts(
    rows: list[dict[str, Any]],
    cols: list[tuple[str, str, float]],
    styles,
    page_width: float,
):
    if not rows:
        return Paragraph("No data available.", styles["SpendlyCell"])

    headers = [Paragraph(label, styles["SpendlyHeader"]) for _, label, _ in cols]
    ratios = [r for _, _, r in cols]
    total = sum(ratios) or 1
    col_widths = [page_width * r / total for r in ratios]

    data = [headers]
    for row in rows:
        data.append([Paragraph(_safe_text(row.get(key)), styles["SpendlyCell"]) for key, _, _ in cols])

    table = LongTable(data, colWidths=col_widths, repeatRows=1, splitByRow=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7.6),
                ("LEADING", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#cbd5e1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.HexColor("#eef2ff")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def _add_pdf_section(story, title: str, rows: list[dict[str, Any]], cols, styles, page_width):
    story.append(Paragraph(title, styles["SpendlySection"]))
    story.append(_pdf_table_from_dicts(rows, cols, styles, page_width))
    story.append(Spacer(1, 8))


def _build_pdf(data: dict[str, Any], user: User) -> bytes:
    styles = _pdf_styles()
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=18,
        rightMargin=18,
        topMargin=18,
        bottomMargin=18,
        title="Spendly Dashboard Export",
        author="Spendly",
        subject="Spendly export",
    )

    story: list[Any] = []
    profile = data["profile"][0] if data["profile"] else {}
    dashboard = data["dashboard"]

    story.append(Paragraph("Spendly Dashboard Export", styles["SpendlyTitle"]))
    story.append(Paragraph(f"Generated on {_fmt_date(_today())}", styles["SpendlySubtle"]))

    profile_rows = [{
        "name": profile.get("full_name") or user.full_name,
        "email": profile.get("email") or user.email,
        "currency": profile.get("currency") or "INR",
        "timezone": profile.get("timezone") or "Asia/Kolkata",
    }]

    summary_rows = [
        {"metric": "Monthly expense", "value": _fmt_money_pdf(dashboard["monthly_expense"])},
        {"metric": "Yearly expense", "value": _fmt_money_pdf(dashboard["yearly_expense"])},
        {"metric": "Monthly income", "value": _fmt_money_pdf(dashboard["monthly_income"])},
        {"metric": "Yearly income", "value": _fmt_money_pdf(dashboard["yearly_income"])},
        {"metric": "Monthly savings", "value": _fmt_money_pdf(dashboard["monthly_savings"])},
        {"metric": "Yearly savings", "value": _fmt_money_pdf(dashboard["yearly_savings"])},
        {"metric": "Lifetime savings", "value": _fmt_money_pdf(dashboard["lifetime_savings"])},
        {"metric": "Expenses", "value": _fmt_number(dashboard["expense_count"])},
        {"metric": "Budgets", "value": _fmt_number(dashboard["budget_count"])},
        {"metric": "Uploaded files", "value": _fmt_number(len(data["uploads"]))},
    ]

    _add_pdf_section(
        story,
        "Profile",
        profile_rows,
        [
            ("name", "Name", 1.3),
            ("email", "Email", 1.8),
            ("currency", "Currency", 0.8),
            ("timezone", "Timezone", 1.1),
        ],
        styles,
        doc.width,
    )

    _add_pdf_section(
        story,
        "Dashboard summary",
        summary_rows,
        [("metric", "Metric", 1.6), ("value", "Value", 1.0)],
        styles,
        doc.width,
    )

    _add_pdf_section(
        story,
        "Budgets",
        data["budgets"],
        [
            ("category_name", "Category", 1.4),
            ("budget_limit", "Budget Limit", 1.0),
            ("spent_amount", "Spent Amount", 1.0),
            ("period", "Period", 0.8),
            ("start_date", "Start Date", 0.9),
            ("end_date", "End Date", 0.9),
            ("notify_percentage", "Notify %", 0.7),
            ("is_active", "Active", 0.6),
            ("created_at", "Created At", 1.1),
        ],
        styles,
        doc.width,
    )

    _add_pdf_section(
        story,
        "Expenses",
        data["expenses"],
        [
            ("vendor_name", "Vendor", 1.4),
            ("category_name", "Category", 1.0),
            ("expense_date", "Expense Date", 0.9),
            ("expense_time", "Expense Time", 0.8),
            ("subtotal", "Subtotal", 0.8),
            ("cgst", "CGST", 0.7),
            ("sgst", "SGST", 0.7),
            ("igst", "IGST", 0.7),
            ("tax_amount", "Tax", 0.8),
            ("discount_amount", "Discount", 0.9),
            ("total_amount", "Total Amount", 0.95),
            ("payment_method", "Payment Method", 0.9),
            ("status", "Status", 1.0),
            ("is_reviewed", "Reviewed", 0.7),
            ("is_manual", "Manual", 0.6),
            ("created_at", "Created At", 1.1),
        ],
        styles,
        doc.width,
    )

    _add_pdf_section(
        story,
        "Expense items",
        data["expense_items"],
        [
            ("item_name", "Item Name", 1.4),
            ("item_category", "Item Category", 0.9),
            ("quantity", "Qty", 0.6),
            ("unit_price", "Unit Price", 0.8),
            ("total_price", "Total Price", 0.8),
            ("vendor_name", "Vendor", 1.1),
            ("expense_date", "Expense Date", 0.9),
            ("created_at", "Created At", 1.0),
        ],
        styles,
        doc.width,
    )

    _add_pdf_section(
        story,
        "Uploaded files",
        data["uploads"],
        [
            ("original_filename", "Original Filename", 2.6),
            ("file_type", "File Type", 0.9),
            ("file_size", "File Size", 0.9),
            ("extraction_status", "Extraction Status", 1.1),
            ("created_at", "Created At", 1.0),
        ],
        styles,
        doc.width,
    )

    _add_pdf_section(
        story,
        "Income",
        data["income"],
        [
            ("source", "Source", 1.6),
            ("amount", "Amount", 0.9),
            ("income_date", "Income Date", 0.9),
            ("notes", "Notes", 2.0),
            ("created_at", "Created At", 1.0),
        ],
        styles,
        doc.width,
    )

    _add_pdf_section(
        story,
        "Top categories",
        dashboard["top_categories"],
        [("name", "Category", 1.6), ("amount", "Amount", 1.0)],
        styles,
        doc.width,
    )

    _add_pdf_section(
        story,
        "Top merchants",
        dashboard["top_merchants"],
        [("name", "Merchant", 1.6), ("amount", "Amount", 1.0)],
        styles,
        doc.width,
    )

    _add_pdf_section(
        story,
        "Top items",
        dashboard["top_items"],
        [("name", "Item", 1.6), ("amount", "Amount", 1.0)],
        styles,
        doc.width,
    )

    _add_pdf_section(
        story,
        "Spending trends",
        dashboard["spending_trends"],
        [("month", "Month", 1.6), ("amount", "Amount", 1.0)],
        styles,
        doc.width,
    )

    doc.build(story, onFirstPage=_page_number, onLaterPages=_page_number)
    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes


# ---------------------------------------------------------------------
# Excel helpers
# ---------------------------------------------------------------------
TITLE_FILL = PatternFill("solid", fgColor="0F172A")
HEADER_FILL = PatternFill("solid", fgColor="111827")
STRIPE_FILL = PatternFill("solid", fgColor="F8FAFC")
BORDER = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)
WHITE_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(color="FFFFFF", bold=True, size=16)
SUBTITLE_FONT = Font(color="D1D5DB", italic=True, size=10)
HEADER_FONT = Font(color="FFFFFF", bold=True)
BODY_FONT = Font(color="111827", size=11)
SECTION_FONT = Font(color="0F172A", bold=True)

EXCEL_MONEY_FMT = '₹#,##0.00'
EXCEL_NUM_FMT = '0.00'
EXCEL_INT_FMT = '0'
EXCEL_DATE_FMT = 'yyyy-mm-dd'
EXCEL_DATETIME_FMT = 'yyyy-mm-dd hh:mm:ss'
EXCEL_TIME_FMT = 'hh:mm:ss'


def _write_sheet(
    ws,
    title: str,
    subtitle: str,
    columns: list[tuple[str, str, str]],
    rows: list[dict[str, Any]],
) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    ws.sheet_view.zoomScale = 95

    last_col = len(columns)
    last_letter = get_column_letter(last_col)

    ws.merge_cells(f"A1:{last_letter}1")
    ws["A1"] = title
    ws["A1"].fill = TITLE_FILL
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells(f"A2:{last_letter}2")
    ws["A2"] = subtitle
    ws["A2"].fill = TITLE_FILL
    ws["A2"].font = SUBTITLE_FONT
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")

    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[4].height = 20

    for col_idx, (_, label, _) in enumerate(columns, start=1):
        cell = ws.cell(row=4, column=col_idx, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    for row_idx, row in enumerate(rows, start=5):
        for col_idx, (key, _, kind) in enumerate(columns, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            value = row.get(key, "")

            if kind == "currency":
                cell.value = _fmt_money_excel(value)
                cell.number_format = EXCEL_MONEY_FMT
                cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
            elif kind == "number":
                cell.value = _fmt_number_excel(value)
                cell.number_format = EXCEL_NUM_FMT
                cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
            elif kind == "integer":
                try:
                    cell.value = int(float(value or 0))
                except Exception:
                    cell.value = 0
                cell.number_format = EXCEL_INT_FMT
                cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
            elif kind == "date":
                cell.value = value
                cell.number_format = EXCEL_DATE_FMT
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            elif kind == "datetime":
                cell.value = value
                cell.number_format = EXCEL_DATETIME_FMT
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            elif kind == "time":
                cell.value = value
                cell.number_format = EXCEL_TIME_FMT
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            elif kind == "bool":
                cell.value = "Yes" if bool(value) else "No"
                cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            else:
                cell.value = _safe_text(value)
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

            cell.font = BODY_FONT
            cell.border = BORDER

        if row_idx % 2 == 0:
            for col_idx in range(1, last_col + 1):
                ws.cell(row=row_idx, column=col_idx).fill = STRIPE_FILL

    ws.auto_filter.ref = f"A4:{last_letter}{max(4, 4 + len(rows))}"

    for col_idx, (_, label, kind) in enumerate(columns, start=1):
        max_len = len(label)
        for row_idx in range(4, 5 + len(rows)):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is None:
                continue
            max_len = max(max_len, len(str(val)))
        if kind in {"currency", "number", "integer"}:
            width = min(max(max_len + 2, 12), 18)
        elif kind in {"date", "datetime", "time"}:
            width = min(max(max_len + 2, 16), 22)
        else:
            width = min(max(max_len + 2, 14), 40)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _build_workbook(data: dict[str, Any], user: User) -> bytes:
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    profile = data["profile"][0] if data["profile"] else {}
    dashboard = data["dashboard"]

    summary_rows: list[dict[str, Any]] = [
        {"section": "Profile", "metric": "Name", "value": profile.get("full_name") or user.full_name, "value_kind": "text"},
        {"section": "Profile", "metric": "Email", "value": profile.get("email") or user.email, "value_kind": "text"},
        {"section": "Profile", "metric": "Currency", "value": profile.get("currency") or "INR", "value_kind": "text"},
        {"section": "Profile", "metric": "Timezone", "value": profile.get("timezone") or "Asia/Kolkata", "value_kind": "text"},
        {"section": "Dashboard Summary", "metric": "Monthly expense", "value": dashboard["monthly_expense"], "value_kind": "currency"},
        {"section": "Dashboard Summary", "metric": "Yearly expense", "value": dashboard["yearly_expense"], "value_kind": "currency"},
        {"section": "Dashboard Summary", "metric": "Monthly income", "value": dashboard["monthly_income"], "value_kind": "currency"},
        {"section": "Dashboard Summary", "metric": "Yearly income", "value": dashboard["yearly_income"], "value_kind": "currency"},
        {"section": "Dashboard Summary", "metric": "Monthly savings", "value": dashboard["monthly_savings"], "value_kind": "currency"},
        {"section": "Dashboard Summary", "metric": "Yearly savings", "value": dashboard["yearly_savings"], "value_kind": "currency"},
        {"section": "Dashboard Summary", "metric": "Lifetime savings", "value": dashboard["lifetime_savings"], "value_kind": "currency"},
        {"section": "Dashboard Summary", "metric": "Expenses", "value": dashboard["expense_count"], "value_kind": "integer"},
        {"section": "Dashboard Summary", "metric": "Budgets", "value": dashboard["budget_count"], "value_kind": "integer"},
        {"section": "Dashboard Summary", "metric": "Uploaded files", "value": len(data["uploads"]), "value_kind": "integer"},
    ]

    ws = wb.create_sheet("Summary")
    _write_sheet(
        ws,
        "Spendly Dashboard Export",
        f"Generated on {_fmt_date(_today())}",
        [
            ("section", "Section", "text"),
            ("metric", "Metric", "text"),
            ("value", "Value", "text"),
        ],
        summary_rows,
    )

    ws = wb.create_sheet("Budgets")
    _write_sheet(
        ws,
        "Budgets",
        "Budget targets and spend tracking",
        [
            ("category_name", "Category", "text"),
            ("budget_limit", "Budget Limit", "currency"),
            ("spent_amount", "Spent Amount", "currency"),
            ("period", "Period", "text"),
            ("start_date", "Start Date", "date"),
            ("end_date", "End Date", "date"),
            ("notify_percentage", "Notify %", "integer"),
            ("is_active", "Active", "bool"),
            ("created_at", "Created At", "datetime"),
        ],
        data["budgets"],
    )

    ws = wb.create_sheet("Expenses")
    _write_sheet(
        ws,
        "Expenses",
        "All saved expenses",
        [
            ("vendor_name", "Vendor", "text"),
            ("category_name", "Category", "text"),
            ("expense_date", "Expense Date", "date"),
            ("expense_time", "Expense Time", "time"),
            ("subtotal", "Subtotal", "currency"),
            ("cgst", "CGST", "currency"),
            ("sgst", "SGST", "currency"),
            ("igst", "IGST", "currency"),
            ("tax_amount", "Tax", "currency"),
            ("discount_amount", "Discount", "currency"),
            ("total_amount", "Total Amount", "currency"),
            ("payment_method", "Payment Method", "text"),
            ("status", "Status", "text"),
            ("is_reviewed", "Reviewed", "bool"),
            ("is_manual", "Manual", "bool"),
            ("created_at", "Created At", "datetime"),
        ],
        data["expenses"],
    )

    ws = wb.create_sheet("Expense Items")
    _write_sheet(
        ws,
        "Expense Items",
        "Line items from your expenses",
        [
            ("item_name", "Item Name", "text"),
            ("item_category", "Item Category", "text"),
            ("quantity", "Qty", "number"),
            ("unit_price", "Unit Price", "currency"),
            ("total_price", "Total Price", "currency"),
            ("vendor_name", "Vendor", "text"),
            ("expense_date", "Expense Date", "date"),
            ("created_at", "Created At", "datetime"),
        ],
        data["expense_items"],
    )

    ws = wb.create_sheet("Uploaded Files")
    _write_sheet(
        ws,
        "Uploaded Files",
        "Only the most useful file details are included",
        [
            ("original_filename", "Original Filename", "text"),
            ("file_type", "File Type", "text"),
            ("file_size", "File Size", "number"),
            ("extraction_status", "Extraction Status", "text"),
            ("created_at", "Created At", "datetime"),
        ],
        data["uploads"],
    )

    ws = wb.create_sheet("Income")
    _write_sheet(
        ws,
        "Income",
        "Income entries recorded in the app",
        [
            ("source", "Source", "text"),
            ("amount", "Amount", "currency"),
            ("income_date", "Income Date", "date"),
            ("notes", "Notes", "text"),
            ("created_at", "Created At", "datetime"),
        ],
        data["income"],
    )

    ws = wb.create_sheet("Top Categories")
    _write_sheet(
        ws,
        "Top Categories",
        "Spending grouped by category",
        [
            ("name", "Category", "text"),
            ("amount", "Amount", "currency"),
        ],
        dashboard["top_categories"],
    )

    ws = wb.create_sheet("Top Merchants")
    _write_sheet(
        ws,
        "Top Merchants",
        "Merchants with the highest spend",
        [
            ("name", "Merchant", "text"),
            ("amount", "Amount", "currency"),
        ],
        dashboard["top_merchants"],
    )

    ws = wb.create_sheet("Top Items")
    _write_sheet(
        ws,
        "Top Items",
        "Items with the highest spend",
        [
            ("name", "Item", "text"),
            ("amount", "Amount", "currency"),
        ],
        dashboard["top_items"],
    )

    ws = wb.create_sheet("Spending Trends")
    _write_sheet(
        ws,
        "Spending Trends",
        "Monthly spend trend",
        [
            ("month", "Month", "text"),
            ("amount", "Amount", "currency"),
        ],
        dashboard["spending_trends"],
    )

    # Move Summary sheet to the front
    wb._sheets = [wb["Summary"]] + [s for s in wb._sheets if s.title != "Summary"]

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


# ---------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------
@router.get("/dashboard/pdf")
def export_dashboard_pdf(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    data = _build_data(db, str(current_user.id))
    pdf_bytes = _build_pdf(data, current_user)
    filename = f"spendly_dashboard_export_{_slugify_filename(current_user.full_name)}_{_today().isoformat()}.pdf"
    return StreamingResponse(
        BytesIO(pdf_bytes),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/dashboard/xlsx")
def export_dashboard_xlsx(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    data = _build_data(db, str(current_user.id))
    xlsx_bytes = _build_workbook(data, current_user)
    filename = f"spendly_dashboard_export_{_slugify_filename(current_user.full_name)}_{_today().isoformat()}.xlsx"
    return StreamingResponse(
        BytesIO(xlsx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )